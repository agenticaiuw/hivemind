import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).parent
DATA = json.loads((BASE / 'purchased-components.json').read_text())
ADD = json.loads((BASE / 'additions-components.json').read_text())['additions']
OUT = Path('/Users/evanliu/agentic-gadget/outputs/01a0d046-7d49-7871-8a87-73ebb2cf093f/Agentic_AI_UW_Component_Costs.docx')
doc = Document()
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.65)
sec.footer_distance = Inches(.28)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(25)
doc.styles['Title'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 1'].font.size = Pt(18)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(9)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(11)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'Agentic AI UW Device Component Costs'
doc.core_properties.author = 'Agentic AI UW'
doc.core_properties.subject = 'Purchased hardware and component pricing for UW Madison ECE review'

def link(p, text, url):
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    r = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
    c = OxmlElement('w:color'); c.set(qn('w:val'), '175E78'); pr.append(c)
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
    r.append(pr); t = OxmlElement('w:t'); t.text = text; r.append(t); h.append(r); p._p.append(h)

def para(text='', bold=None):
    after_table = len(doc._element.body) > 1 and doc._element.body[-2].tag == qn('w:tbl')
    p = doc.add_paragraph()
    if after_table and text:
        p.paragraph_format.space_before = Pt(7)
    if bold and text.startswith(bold):
        p.add_run(bold).bold = True; p.add_run(text[len(bold):])
    else: p.add_run(text)
    return p

def table(headers, widths, rows, size=10):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, widths): col.width = Inches(w)
    for c, h in zip(t.rows[0].cells, headers): c.text = h
    rep = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells = t.add_row().cells
        for c, val in zip(cells, row):
            p = c.paragraphs[0]
            if isinstance(val, list):
                for i, chunk in enumerate(val):
                    if i: p.add_run('\n')
                    if isinstance(chunk, tuple): link(p, *chunk)
                    else: p.add_run(chunk)
            else: p.add_run(str(val))
    for ri, row in enumerate(t.rows):
        trpr = row._tr.get_or_add_trPr(); ns = OxmlElement('w:cantSplit'); trpr.append(ns)
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b = OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
            tcpr.append(borders)
            mar = OxmlElement('w:tcMar')
            for side, value in [('top',75),('bottom',75),('left',90),('right',90)]:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), str(value)); e.set(qn('w:type'),'dxa'); mar.append(e)
            tcpr.append(mar)
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF1'); tcpr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                if ci == 0 and len(headers)>3: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(size); r.bold = ri == 0
                # Hyperlink runs do not appear in python-docx paragraph.runs.
                for rpr in p._p.xpath('.//w:hyperlink/w:r/w:rPr'):
                    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size*2))); rpr.append(sz)
    return t

f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run('Agentic AI UW  |  '); r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)

doc.add_paragraph('Agentic AI UW Device Component Costs', 'Title')
para('UW–Madison ECE • Prices checked 23 September 2026 • Single-unit component prices in USD')
para('The final product uses a Bluetooth module. The ESP32 is testing hardware only. Product names link directly to suppliers. Existing purchases and added features are combined below.')
doc.add_heading('1 Breakout boards and prototype hardware',1)
short = ['MAX98357A I2S 3 W amplifier breakout','PDM MEMS microphone breakout','DRV2605L haptic driver breakout','Vibrating mini motor disc','LiPo battery 3.7 V 500 mAh','Mini oval speaker 8 Ω 1 W','LSM6DSOX six-axis IMU breakout','Nordic nRF9160 Development Kit','HUZZAH32 ESP32 Feather — testing only','microSD breakout 5 V or 3 V','DFRobot FIT0642 64 GB microSD card']
functions = ['Amplifier','Microphone','Haptic driver','Vibration motor','Battery','Speaker','Motion sensor','LTE-M testing','ESP32 testing','Card interface','Storage']
rows=[]
for i,(p,n,f) in enumerate(zip(DATA['purchased'],short,functions),1):
    rows.append([i,f,[(n,p['product_url'])],f"${p['paid_unit']:.2f}"])
rows += [
    [12,'Bluetooth',[('Raytac MDBT53-DB-40 nRF5340 board','https://www.digikey.com/en/products/detail/raytac/MDBT53-DB-40/16630695')],'$17.50'],
    [13,'Camera',[('Arducam B006701 SPI camera',ADD[0]['board_url'])],'$25.99'],
    [14,'Fingerprint',[('SparkFun FPC2534 Qwiic',ADD[1]['board_url'])],'$59.95'],
    [15,'Solar panel',[('Voltaic P122 panel',ADD[2]['board_url'])],'$5.95'],
    [16,'Solar charger',[('Adafruit BQ24074 charger',ADD[3]['board_url'])],'$14.95'],
]
table(['Item','Function','Product','Cost'],[.42,1.18,4.55,1.05],rows,10)
para('Items 1–11 use the actual receipt prices; the remaining entries use current single-unit listings. Prices exclude tax, shipping, and tariffs. Cellular and payment hardware are listed as individual ICs in the final-parts table on page 2.')
para('The TDK microphone and DVP camera need interface validation. BLE does not provide Classic Bluetooth/A2DP. Payment provisioning, carrier service, support circuitry, and assembly are additional requirements; these tables are not a complete assembled-device quote.')

doc.add_page_break()
doc.add_heading('2 Final chips components and PCB purchase',1)
para('The Bluetooth module replaces the ESP32. Prices are per unit unless stated. Rows 12 and 12A are alternatives, not additive. Cellular options 8 and 17 have different calling capabilities. Modules and complete physical parts are labeled explicitly.')
B = {x['purchased_id']:x for x in DATA['bare'] if x['type'] not in ['bare microphone alternative', 'bare SoC']}
def component(k,label):return [(label,B[k]['url'])]
mic_alt=next(x for x in DATA['bare'] if x['purchased_id']=='mic' and x['unit_price'] is not None)
rows=[
    [1,'Amplifier',component('amp','MAX98357AETE+T IC'),'$4.08'],
    [2,'Microphone',[('TDK MMICT5838 PDM mic — alternative',mic_alt['url'])],'$2.67'],
    [3,'Haptic driver',component('haptic_driver','DRV2605LDGSR IC'),'$2.10'],
    [4,'Vibration motor',component('motor','Mini motor disc — complete actuator'),'$1.95'],
    [5,'Battery',component('battery','3.7 V 500 mAh LiPo pack'),'$7.95'],
    [6,'Speaker',component('speaker','8 Ω 1 W oval speaker'),'$1.95'],
    [7,'Motion sensor',component('imu','LSM6DSOXTR IC — backorder'),'$5.30'],
    [8,'Cellular data',[('NRF9160-SICA-B1A-R7 SiP — data/SMS',B['cellular']['url'])],'$31.56'],
    [10,'Card interface',component('sd_breakout','GCT MEM2090 microSD socket'),'$1.50'],
    [11,'Storage',component('sd_card','DFRobot FIT0642 64 GB card — backorder'),'$18.90'],
    [12,'Bluetooth',[('Raytac MDBT53-P1M nRF5340 module','https://www.digikey.com/en/products/detail/raytac/MDBT53-P1M/16603007')],'$7.00'],
    ['12A','Bare alternative',[('Nordic NRF5340-CLAA-R SoC','https://www.digikey.com/en/products/detail/nordic-semiconductor-asa/NRF5340-CLAA-R/14323741')],'$8.65'],
    [13,'Camera',[('OV2640 M0031 lens and FPC assembly',ADD[0]['bare_url'])],'$6.99'],
    [14,'Fingerprint',[('FPC2534AP LGA sensor system',ADD[1]['bare_url'])],'$23.85'],
    [15,'Solar panel',[('Voltaic P122 finished panel',ADD[2]['board_url'])],'$5.95'],
    [16,'Solar charger',[('TI BQ24074RGTR IC',ADD[3]['bare_url'])],'$2.43'],
    [17,'Voice modem IC',[('Qualcomm MDM9207 bare IC — see note','https://tech-electr.com/product/mdm9207-qualcomm-genuine-reliable-electronic-components/')],'$9.20'],
    [18,'NFC / secure IC',[('NXP SN100U / 100VB27 — repair-market IC','https://www.phonelcdparts.com/iphone-xr-xs-xs-max-ipad-pro-11-1st-gen-2018-ipad-pro-11-2rd-gen-2020-ipad-7-2019-ic-100vb27-xr')],'$2.14'],
    [19,'PCB fabrication',[('4-layer PCB • 32 mm • 0.8 mm thick','https://jlcpcb.com/resources/pcb-thickness')],'~$8 each\n~$240 / 30'],
]
table(['Item','Function','Product','Cost'],[.42,1.18,4.55,1.05],rows,10)
para('Nordic options: the $7 Bluetooth module contains an nRF5340. Its bare-SoC alternative is $8.65 for one, plus RF/support circuitry. The $31.56 nRF9160 SiP supports data/SMS, not VoLTE calls. ESP32 remains testing-only.',bold='Nordic options:')
p=para('PCB: $240 is a planning allowance for 30 bare 0.8 mm boards, excluding components, soldering/assembly, shipping, and tax. For comparison, ');link(p,'OSH Park’s published four-layer rate','https://docs.oshpark.com/services/four-layer/');p.add_run(' gives about $158.72 for 30 at a 32 × 32 mm bounding size, but at 1.6 mm thickness. Final cost needs routed Gerbers and a supplier quote.')
p=para('Bare modem: $9.20 buys one MDM9207 IC from an independent seller. ');link(p,'Qualcomm specifies VoLTE for MDM9207-1','https://s204.q4cdn.com/645488518/files/doc_news/2015/10/2015-10-26_Qualcomm_Announces_New_Modem_Solutions_Designed_725.pdf');p.add_run('; the listing does not confirm that variant. RF transceiver, power IC, memory, antenna, firmware and carrier compatibility require separate validation/costing.')
p=para('Payment IC: ');link(p,'TechInsights identifies 100VB27 as SN100U','https://www.techinsights.com/blog/google-pixel-3-xl-teardown');p.add_run(', integrating NFC, a secure element and eSIM. The $2.14 repair-market listing covers one physical chip; keys, payment provisioning and custom-board usability are unverified. It is not a ready-to-pay wallet. Both bare IC prices exclude shipping/tax and are sourcing leads, not validated design selections.')
OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
print(OUT)
